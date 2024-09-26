import openpyxl


# Fonction pour ajouter du texte à une cellule avec délimiteur

def add_text_to_cell(file_path: str, cell_ref: str, text: str, delimiter: str = "", sheets: list = None) -> None:
    """
    Ajoute du texte à une cellule spécifique, avec un délimiteur si la cellule contient déjà du texte.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_ref (str): La référence de la cellule (ex: 'A1').
        text (str): Le texte à ajouter à la cellule.
        delimiter (str): Le délimiteur à utiliser entre l'ancien texte et le nouveau (par défaut, pas de délimiteur).
        sheets (list): La liste des feuilles où effectuer la modification (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell = sheet[cell_ref]
        if cell.value:  # Vérifie si la cellule contient déjà du texte
            cell.value = f"{cell.value}{delimiter}{text}"
        else:
            cell.value = text

    wb.save(file_path)
    print(f"Le texte a été ajouté à la cellule {cell_ref} avec le délimiteur '{delimiter}' dans les feuilles {sheet_names}.")


# Fonction pour ajouter du texte à une colonne avec délimiteur

def add_text_to_column(file_path: str, column_header: str, text: str, delimiter: str = "", sheets: list = None) -> None:
    """
    Ajoute du texte à toutes les cellules d'une colonne spécifiée, avec un délimiteur si la cellule contient déjà du texte.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        column_header (str): L'entête de la colonne à modifier.
        text (str): Le texte à ajouter à chaque cellule de la colonne.
        delimiter (str): Le délimiteur à utiliser entre l'ancien texte et le nouveau (par défaut, pas de délimiteur).
        sheets (list): La liste des feuilles où effectuer la modification (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column_index)
            if cell.value:  # Vérifie si la cellule contient déjà du texte
                cell.value = f"{cell.value}{delimiter}{text}"
            else:
                cell.value = text

    wb.save(file_path)
    print(f"Le texte a été ajouté à toutes les cellules de la colonne '{column_header}' avec le délimiteur '{delimiter}' dans les feuilles {sheet_names}.")


# Fonction pour ajouter du texte à une ligne avec délimiteur

def add_text_to_row(file_path: str, row_number: int, text: str, delimiter: str = "", sheets: list = None) -> None:
    """
    Ajoute du texte à toutes les cellules d'une ligne spécifiée, avec un délimiteur si la cellule contient déjà du texte.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        row_number (int): Le numéro de la ligne à modifier.
        text (str): Le texte à ajouter à chaque cellule de la ligne.
        delimiter (str): Le délimiteur à utiliser entre l'ancien texte et le nouveau (par défaut, pas de délimiteur).
        sheets (list): La liste des feuilles où effectuer la modification (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]

        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_number, column=column)
            if cell.value:  # Vérifie si la cellule contient déjà du texte
                cell.value = f"{cell.value}{delimiter}{text}"
            else:
                cell.value = text

    wb.save(file_path)
    print(f"Le texte a été ajouté à toutes les cellules de la ligne {row_number} avec le délimiteur '{delimiter}' dans les feuilles {sheet_names}.")


# Fonction utilitaire pour trouver l'index d'une colonne à partir de l'entête

def get_column_index(sheet, column_header: str) -> int:
    """Retourne l'index de la colonne pour l'entête donné."""
    for cell in sheet[1]:
        if cell.value == column_header:
            return cell.column
    raise ValueError(f"Colonne avec l'entête '{column_header}' non trouvée.")
