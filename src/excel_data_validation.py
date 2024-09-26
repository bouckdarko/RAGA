import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation


def apply_list_validation(file_path: str, cell_range: str, values: list, sheets: list = None) -> None:
    """
    Applique une validation de type liste à une plage de cellules, permettant seulement les valeurs spécifiées.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_range (str): La plage de cellules à laquelle la validation est appliquée (ex: 'A2:A10').
        values (list): La liste des valeurs autorisées.
        sheets (list): La liste des feuilles où effectuer la validation (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]
    values_str = ",".join(map(str, values))

    # Création de la validation
    validation = DataValidation(type="list", formula1=f'"{values_str}"', allow_blank=True)

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        sheet.add_data_validation(validation)
        validation.add(sheet[cell_range])

    wb.save(file_path)
    print(f"Validation de liste appliquée aux cellules {cell_range} avec les valeurs {values_str} dans les feuilles {sheet_names}.")


def apply_numeric_validation(file_path: str, cell_range: str, min_value: float, max_value: float, sheets: list = None) -> None:
    """
    Applique une validation numérique à une plage de cellules, restreignant les valeurs entre min_value et max_value.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_range (str): La plage de cellules à laquelle la validation est appliquée (ex: 'B2:B10').
        min_value (float): La valeur minimale autorisée.
        max_value (float): La valeur maximale autorisée.
        sheets (list): La liste des feuilles où effectuer la validation (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    # Création de la validation
    validation = DataValidation(type="decimal", operator="between", formula1=str(min_value), formula2=str(max_value), allow_blank=True)

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        sheet.add_data_validation(validation)
        validation.add(sheet[cell_range])

    wb.save(file_path)
    print(f"Validation numérique appliquée aux cellules {cell_range} (min: {min_value}, max: {max_value}) dans les feuilles {sheet_names}.")


def apply_date_validation(file_path: str, cell_range: str, start_date: str, end_date: str, sheets: list = None) -> None:
    """
    Applique une validation de date à une plage de cellules, restreignant les dates entre start_date et end_date.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_range (str): La plage de cellules à laquelle la validation est appliquée (ex: 'C2:C10').
        start_date (str): La date minimale autorisée (au format 'YYYY-MM-DD').
        end_date (str): La date maximale autorisée (au format 'YYYY-MM-DD').
        sheets (list): La liste des feuilles où effectuer la validation (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    # Création de la validation
    validation = DataValidation(type="date", operator="between", formula1=f'"{start_date}"', formula2=f'"{end_date}"', allow_blank=True)

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        sheet.add_data_validation(validation)
        validation.add(sheet[cell_range])

    wb.save(file_path)
    print(f"Validation de date appliquée aux cellules {cell_range} (de {start_date} à {end_date}) dans les feuilles {sheet_names}.")


def apply_text_length_validation(file_path: str, cell_range: str, max_length: int, sheets: list = None) -> None:
    """
    Applique une validation de longueur de texte à une plage de cellules, restreignant la longueur maximale.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_range (str): La plage de cellules à laquelle la validation est appliquée (ex: 'D2:D10').
        max_length (int): La longueur maximale autorisée.
        sheets (list): La liste des feuilles où effectuer la validation (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    # Création de la validation
    validation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(max_length), allow_blank=True)

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        sheet.add_data_validation(validation)
        validation.add(sheet[cell_range])

    wb.save(file_path)
    print(f"Validation de longueur de texte appliquée aux cellules {cell_range} (max: {max_length} caractères) dans les feuilles {sheet_names}.")

