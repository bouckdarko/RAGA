import openpyxl


# Conversion au format "standard" (General)

def set_column_to_standard(file_path: str, column_header: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column_index)
            cell.number_format = 'General'

    wb.save(file_path)
    print(f"Les cellules de la colonne '{column_header}' dans les feuilles {sheet_names} ont été mises au format 'standard'.")


def set_row_to_standard(file_path: str, row_number: int, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]

        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_number, column=column)
            cell.number_format = 'General'

    wb.save(file_path)
    print(f"Les cellules de la ligne {row_number} dans les feuilles {sheet_names} ont été mises au format 'standard'.")


def set_cell_to_standard(file_path: str, cell_ref: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell = sheet[cell_ref]
        cell.number_format = 'General'

    wb.save(file_path)
    print(f"La cellule {cell_ref} dans les feuilles {sheet_names} a été mise au format 'standard'.")


# Conversion au format "date"

def set_column_to_date(file_path: str, column_header: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column_index)
            cell.number_format = 'YYYY-MM-DD'

    wb.save(file_path)
    print(f"Les cellules de la colonne '{column_header}' dans les feuilles {sheet_names} ont été mises au format 'date'.")


def set_row_to_date(file_path: str, row_number: int, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]

        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_number, column=column)
            cell.number_format = 'YYYY-MM-DD'

    wb.save(file_path)
    print(f"Les cellules de la ligne {row_number} dans les feuilles {sheet_names} ont été mises au format 'date'.")


def set_cell_to_date(file_path: str, cell_ref: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell = sheet[cell_ref]
        cell.number_format = 'YYYY-MM-DD'

    wb.save(file_path)
    print(f"La cellule {cell_ref} dans les feuilles {sheet_names} a été mise au format 'date'.")


# Conversion au format "numérique"

def set_column_to_numeric(file_path: str, column_header: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column_index)
            cell.number_format = '0.00'  # Format numérique avec 2 décimales

    wb.save(file_path)
    print(f"Les cellules de la colonne '{column_header}' dans les feuilles {sheet_names} ont été mises au format 'numérique'.")


def set_row_to_numeric(file_path: str, row_number: int, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]

        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_number, column=column)
            cell.number_format = '0.00'  # Format numérique avec 2 décimales

    wb.save(file_path)
    print(f"Les cellules de la ligne {row_number} dans les feuilles {sheet_names} ont été mises au format 'numérique'.")


def set_cell_to_numeric(file_path: str, cell_ref: str, sheets: list = None) -> None:
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell = sheet[cell_ref]
        cell.number_format = '0.00'  # Format numérique avec 2 décimales

    wb.save(file_path)
    print(f"La cellule {cell_ref} dans les feuilles {sheet_names} a été mise au format 'numérique'.")


# Fonction utilitaire pour trouver l'index d'une colonne à partir de l'entête

def get_column_index(sheet, column_header: str) -> int:
    for cell in sheet[1]:
        if cell.value == column_header:
            return cell.column
    raise ValueError(f"Colonne avec l'entête '{column_header}' non trouvée.")
