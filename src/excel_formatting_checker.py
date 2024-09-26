import openpyxl
import re
from openpyxl.utils import get_column_letter
from datetime import datetime


# Vérification des colonnes

def check_column_format(file_path: str, column_header: str = "TRAD", sheets: list = None) -> None:
    """
    Vérifie que toutes les cellules de la colonne spécifiée par l'entête
    respectent le format "TR" suivi de 13 chiffres.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        # Expression régulière pour vérifier le format "TR" suivi de 13 chiffres
        pattern = re.compile(r'^TR\d{13}$')
        invalid_cells = []

        for row in range(2, sheet.max_row + 1):
            cell_value = sheet.cell(row=row, column=column_index).value
            if not pattern.match(str(cell_value)):
                cell_ref = sheet.cell(row=row, column=column_index).coordinate
                invalid_cells.append(f"{sheet_name}!{cell_ref}")

        display_results(invalid_cells, "Toutes les cellules respectent le format 'TR' suivi de 13 chiffres.")


def check_numeric_column(file_path: str, column_header: str = "pouet", sheets: list = None) -> None:
    """
    Vérifie que toutes les cellules de la colonne spécifiée contiennent
    uniquement des valeurs numériques.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)
        invalid_cells = []

        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column_index)
            cell_value = cell.value

            if not is_numeric(cell_value):
                cell_ref = cell.coordinate
                invalid_cells.append(f"{sheet_name}!{cell_ref}")

        display_results(invalid_cells, "Toutes les cellules contiennent des valeurs numériques.")


def check_date_column(file_path: str, column_header: str = "Date", sheets: list = None) -> None:
    """
    Vérifie que toutes les cellules de la colonne spécifiée contiennent
    uniquement des dates valides.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)
        invalid_cells = []

        for row in range(2, sheet.max_row + 1):
            cell_value = sheet.cell(row=row, column=column_index).value

            if not is_valid_date(cell_value):
                cell_ref = sheet.cell(row=row, column=column_index).coordinate
                invalid_cells.append(f"{sheet_name}!{cell_ref}")

        display_results(invalid_cells, "Toutes les cellules contiennent des dates valides.")


# Vérification des lignes

def check_row_format(file_path: str, row_number: int, column_header: str = "TRAD", sheets: list = None) -> None:
    """
    Vérifie que la cellule de la ligne spécifiée par l'entête respecte
    le format "TR" suivi de 13 chiffres.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]
    pattern = re.compile(r'^TR\d{13}$')

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        cell_value = sheet.cell(row=row_number, column=column_index).value
        if not pattern.match(str(cell_value)):
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} n'a pas le format attendu.")
        else:
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} respecte le format 'TR' suivi de 13 chiffres.")


def check_numeric_row(file_path: str, row_number: int, column_header: str = "pouet", sheets: list = None) -> None:
    """
    Vérifie que la cellule de la ligne spécifiée contient uniquement une valeur numérique.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        cell_value = sheet.cell(row=row_number, column=column_index).value
        if not is_numeric(cell_value):
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} ne contient pas de valeur numérique.")
        else:
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} contient une valeur numérique.")


def check_date_row(file_path: str, row_number: int, column_header: str = "Date", sheets: list = None) -> None:
    """
    Vérifie que la cellule de la ligne spécifiée contient une date valide.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        column_index = get_column_index(sheet, column_header)

        cell_value = sheet.cell(row=row_number, column=column_index).value
        if not is_valid_date(cell_value):
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} ne contient pas une date valide.")
        else:
            print(f"La cellule {sheet_name}!{sheet.cell(row=row_number, column=column_index).coordinate} contient une date valide.")


# Vérification des cellules spécifiques

def check_cell_format(file_path: str, cell_ref: str, sheets: list = None) -> None:
    """
    Vérifie qu'une cellule spécifique respecte le format "TR" suivi de 13 chiffres.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]
    pattern = re.compile(r'^TR\d{13}$')

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell_value = sheet[cell_ref].value
        if not pattern.match(str(cell_value)):
            print(f"La cellule {sheet_name}!{cell_ref} n'a pas le format attendu.")
        else:
            print(f"La cellule {sheet_name}!{cell_ref} respecte le format 'TR' suivi de 13 chiffres.")


def check_numeric_cell(file_path: str, cell_ref: str, sheets: list = None) -> None:
    """
    Vérifie qu'une cellule spécifique contient uniquement une valeur numérique.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell_value = sheet[cell_ref].value
        if not is_numeric(cell_value):
            print(f"La cellule {sheet_name}!{cell_ref} ne contient pas de valeur numérique.")
        else:
            print(f"La cellule {sheet_name}!{cell_ref} contient une valeur numérique.")


def check_date_cell(file_path: str, cell_ref: str, sheets: list = None) -> None:
    """
    Vérifie qu'une cellule spécifique contient une date valide.
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        cell_value = sheet[cell_ref].value
        if not is_valid_date(cell_value):
            print(f"La cellule {sheet_name}!{cell_ref} ne contient pas une date valide.")
        else:
            print(f"La cellule {sheet_name}!{cell_ref} contient une date valide.")


# Fonctions utilitaires

def get_column_index(sheet, column_header: str) -> int:
    """Retourne l'index de la colonne pour l'entête donné."""
    for cell in sheet[1]:
        if cell.value == column_header:
            return cell.column
    raise ValueError(f"Colonne avec l'entête '{column_header}' non trouvée.")


def is_numeric(value) -> bool:
    """Vérifie si une valeur est numérique."""
    try:
        float(value)
        return True
    except (ValueError, TypeError):
        return False


def is_valid_date(value) -> bool:
    """Vérifie si une valeur est une date valide."""
    if isinstance(value, datetime):
        return True
    try:
        datetime.strptime(str(value), "%Y-%m-%d")
        return True
    except (ValueError, TypeError):
        return False


def display_results(invalid_cells, success_message: str) -> None:
    """Affiche les résultats des vérifications."""
    if invalid_cells:
        print(f"Les cellules {', '.join(invalid_cells)} ne respectent pas le format attendu.")
    else:
        print(success_message)
