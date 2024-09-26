"""
Module excel_alignment_manager

Ce module permet de gérer l'alignement du texte dans les cellules d'un fichier Excel
en utilisant openpyxl. Il offre des fonctionnalités pour aligner le texte verticalement
(haut, milieu, bas), horizontalement (gauche, centre, droite), et permet également
de renvoyer le texte à la ligne automatiquement.

Le module inclut également une fonction pour fusionner des cellules.

Fonctionnalités :
----------------
1. Aligner le texte dans les cellules (haut, milieu, bas, gauche, centre, droite).
2. Renvoyer le texte à la ligne automatiquement.
3. Fusionner des cellules.

Auteur : Bouck DARKO
Version : 1.0
"""

import openpyxl
from openpyxl.styles import Alignment


def set_alignment(file_path: str, cell_range: str, horizontal: str = None, vertical: str = None, wrap_text: bool = False, sheets: list = None) -> None:
    """
    Définit l'alignement du texte dans une plage de cellules.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        cell_range (str): La plage de cellules à laquelle l'alignement est appliqué (ex: 'A1:C3').
        horizontal (str): L'alignement horizontal ('left', 'center', 'right').
        vertical (str): L'alignement vertical ('top', 'center', 'bottom').
        wrap_text (bool): Si True, renvoie le texte à la ligne automatiquement.
        sheets (list): La liste des feuilles où appliquer l'alignement (par défaut, feuille active).
    """
    wb = openpyxl.load_workbook(file_path)
    sheet_names = sheets if sheets else [wb.active.title]

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)

        # Appliquer l'alignement à la plage de cellules
        for row in sheet[cell_range]:
            for cell in row:
                cell.alignment = alignment

    wb.save(file_path)
    print(f"L'alignement a été appliqué à la plage {cell_range} dans les feuilles {sheet_names}.")


def merge_cells(file_path: str, sheet_name: str, start_cell: str, end_cell: str) -> None:
    """
    Fusionne une plage de cellules dans une feuille spécifiée.

    Args:
        file_path (str): Le chemin vers le fichier Excel.
        sheet_name (str): Le nom de la feuille où effectuer la fusion.
        start_cell (str): La cellule de début de la fusion (ex: 'A1').
        end_cell (str): La cellule de fin de la fusion (ex: 'C3').
    """
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]

    # Fusionner les cellules spécifiées
    sheet.merge_cells(f"{start_cell}:{end_cell}")

    wb.save(file_path)
    print(f"Les cellules de {start_cell} à {end_cell} ont été fusionnées dans la feuille '{sheet_name}'.")
