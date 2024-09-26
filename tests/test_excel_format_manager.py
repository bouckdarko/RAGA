import pytest
import openpyxl
from excel_format_tools import (
    set_column_to_standard, set_row_to_standard, set_cell_to_standard,
    set_column_to_date, set_row_to_date, set_cell_to_date,
    set_column_to_numeric, set_row_to_numeric, set_cell_to_numeric
)

# Fixtures pour créer des fichiers Excel temporaires pour les tests
@pytest.fixture
def excel_file(tmp_path):
    """Crée un fichier Excel temporaire avec différentes feuilles pour les tests."""
    file_path = tmp_path / "test_workbook.xlsx"
    wb = openpyxl.Workbook()

    # Feuille 1
    sheet1 = wb.active
    sheet1.title = "Feuille1"
    sheet1.append(["TRRA", "pouet", "Date"])
    sheet1.append(["TR1234567890123", 123, "2023-09-23"])
    sheet1.append(["TR0987654321098", 456, "2024-01-01"])

    # Feuille 2
    sheet2 = wb.create_sheet(title="Feuille2")
    sheet2.append(["TRRA", "pouet", "Date"])
    sheet2.append(["TR9999999999999", 789, "2025-05-05"])

    wb.save(file_path)
    return file_path

# Helper function to check the number format of a cell
def check_number_format(sheet, cell_ref, expected_format):
    """Vérifie le format numérique de la cellule spécifiée."""
    return sheet[cell_ref].number_format == expected_format

# Tests pour les conversions au format "standard"

def test_set_column_to_standard(excel_file):
    """Test pour mettre une colonne entière au format 'standard'."""
    set_column_to_standard(excel_file, "TRRA", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "A2", "General")
    assert check_number_format(sheet, "A3", "General")


def test_set_row_to_standard(excel_file):
    """Test pour mettre une ligne entière au format 'standard'."""
    set_row_to_standard(excel_file, 2, sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "A2", "General")
    assert check_number_format(sheet, "B2", "General")
    assert check_number_format(sheet, "C2", "General")


def test_set_cell_to_standard(excel_file):
    """Test pour mettre une cellule spécifique au format 'standard'."""
    set_cell_to_standard(excel_file, "A2", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "A2", "General")


# Tests pour les conversions au format "date"

def test_set_column_to_date(excel_file):
    """Test pour mettre une colonne entière au format 'date'."""
    set_column_to_date(excel_file, "Date", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "C2", "YYYY-MM-DD")
    assert check_number_format(sheet, "C3", "YYYY-MM-DD")


def test_set_row_to_date(excel_file):
    """Test pour mettre une ligne entière au format 'date'."""
    set_row_to_date(excel_file, 2, sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "A2", "YYYY-MM-DD")
    assert check_number_format(sheet, "B2", "YYYY-MM-DD")
    assert check_number_format(sheet, "C2", "YYYY-MM-DD")


def test_set_cell_to_date(excel_file):
    """Test pour mettre une cellule spécifique au format 'date'."""
    set_cell_to_date(excel_file, "C2", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "C2", "YYYY-MM-DD")


# Tests pour les conversions au format "numérique"

def test_set_column_to_numeric(excel_file):
    """Test pour mettre une colonne entière au format 'numérique'."""
    set_column_to_numeric(excel_file, "pouet", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "B2", "0.00")
    assert check_number_format(sheet, "B3", "0.00")


def test_set_row_to_numeric(excel_file):
    """Test pour mettre une ligne entière au format 'numérique'."""
    set_row_to_numeric(excel_file, 2, sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "A2", "0.00")
    assert check_number_format(sheet, "B2", "0.00")
    assert check_number_format(sheet, "C2", "0.00")


def test_set_cell_to_numeric(excel_file):
    """Test pour mettre une cellule spécifique au format 'numérique'."""
    set_cell_to_numeric(excel_file, "B2", sheets=["Feuille1"])
    
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    assert check_number_format(sheet, "B2", "0.00")
