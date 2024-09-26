import pytest
import openpyxl
from openpyxl.styles import PatternFill, Font
from excel_color_manager import (
    apply_color_to_cell, apply_color_to_column, apply_color_to_row, apply_color_to_range,
    apply_text_color_to_column, apply_text_color_to_row, apply_text_color_to_range,
    apply_text_format_to_column, apply_text_format_to_row, apply_text_format_to_range
)


@pytest.fixture
def excel_file(tmp_path):
    """Crée un fichier Excel temporaire pour les tests."""
    file_path = tmp_path / "test_color_manager.xlsx"
    wb = openpyxl.Workbook()

    # Feuille de test
    sheet = wb.active
    sheet.title = "Feuille1"
    sheet.append(["Col1", "Col2", "Col3", "Col4", "Col5"])
    for i in range(1, 6):
        sheet.append([f"Row{i}Col1", f"Row{i}Col2", f"Row{i}Col3", f"Row{i}Col4", f"Row{i}Col5"])

    wb.save(file_path)
    return file_path


def get_cell_fill(file_path, sheet_name, cell_ref):
    """Récupère la couleur de fond d'une cellule."""
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]
    return sheet[cell_ref].fill


def get_cell_font_color(file_path, sheet_name, cell_ref):
    """Récupère la couleur du texte d'une cellule."""
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]
    return sheet[cell_ref].font.color


def test_apply_color_to_cell(excel_file):
    """Test pour appliquer une couleur de fond à une cellule spécifique."""
    apply_color_to_cell(excel_file, "A1", "FFFF00")
    fill = get_cell_fill(excel_file, "Feuille1", "A1")
    assert fill.start_color.rgb == "00FFFF00"  # Couleur jaune


def test_apply_color_to_column(excel_file):
    """Test pour appliquer une couleur de fond à une colonne."""
    apply_color_to_column(excel_file, "Col2", "00FF00", only_non_empty=True)
    fill = get_cell_fill(excel_file, "Feuille1", "B2")
    assert fill.start_color.rgb == "0000FF00"  # Couleur verte


def test_apply_color_to_row(excel_file):
    """Test pour appliquer une couleur de fond à une ligne."""
    apply_color_to_row(excel_file, 2, "0000FF", only_non_empty=True)
    fill = get_cell_fill(excel_file, "Feuille1", "A2")
    assert fill.start_color.rgb == "000000FF"  # Couleur bleue


def test_apply_color_to_range(excel_file):
    """Test pour appliquer une couleur de fond à une plage de cellules."""
    apply_color_to_range(excel_file, "A", "C", "FF00FF", 1, 3, only_non_empty=True)
    fill = get_cell_fill(excel_file, "Feuille1", "A2")
    assert fill.start_color.rgb == "00FF00FF"  # Couleur magenta


def test_apply_text_color_to_column(excel_file):
    """Test pour appliquer une couleur au texte dans une colonne."""
    apply_text_color_to_column(excel_file, "Col2", "FF0000")
    font_color = get_cell_font_color(excel_file, "Feuille1", "B2")
    assert font_color.rgb == "00FF0000"  # Texte rouge


def test_apply_text_color_to_row(excel_file):
    """Test pour appliquer une couleur au texte dans une ligne."""
    apply_text_color_to_row(excel_file, 2, "0000FF")
    font_color = get_cell_font_color(excel_file, "Feuille1", "A2")
    assert font_color.rgb == "000000FF"  # Texte bleu


def test_apply_text_color_to_range(excel_file):
    """Test pour appliquer une couleur au texte dans une plage de cellules."""
    apply_text_color_to_range(excel_file, "A", "C", "00FF00", 1, 3)
    font_color = get_cell_font_color(excel_file, "Feuille1", "A2")
    assert font_color.rgb == "0000FF00"  # Texte vert


def test_apply_text_format_to_column(excel_file):
    """Test pour appliquer des formats de texte à une colonne."""
    apply_text_format_to_column(excel_file, "Col2", bold=True, italic=True)
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    cell = sheet["B2"]
    assert cell.font.bold is True
    assert cell.font.italic is True


def test_apply_text_format_to_row(excel_file):
    """Test pour appliquer des formats de texte à une ligne."""
    apply_text_format_to_row(excel_file, 2, underline=True, strike=True)
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    cell = sheet["A2"]
    assert cell.font.underline == "single"
    assert cell.font.strike is True


def test_apply_text_format_to_range(excel_file):
    """Test pour appliquer des formats de texte à une plage de cellules."""
    apply_text_format_to_range(excel_file, "A", "C", bold=True, italic=True, start_row=1, end_row=3)
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]
    cell = sheet["A2"]
    assert cell.font.bold is True
    assert cell.font.italic is True
