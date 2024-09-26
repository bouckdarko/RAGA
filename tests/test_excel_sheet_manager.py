import pytest
import openpyxl
from excel_sheet_manager import (
    create_sheet, delete_sheet, set_active_sheet, change_sheet_color, 
    lock_sheet, unlock_sheet, copy_sheet, reorder_sheets, get_sheet_list, get_sheet_dict
)


@pytest.fixture
def excel_file(tmp_path):
    """Crée un fichier Excel temporaire pour les tests."""
    file_path = tmp_path / "test_workbook.xlsx"
    wb = openpyxl.Workbook()

    # Ajouter des feuilles de test
    sheet1 = wb.active
    sheet1.title = "Feuille1"
    wb.create_sheet("Feuille2")
    wb.create_sheet("Feuille3")

    wb.save(file_path)
    return file_path


def test_create_sheet(excel_file):
    """Test pour créer une nouvelle feuille."""
    create_sheet(excel_file, "NouvelleFeuille")
    wb = openpyxl.load_workbook(excel_file)
    assert "NouvelleFeuille" in wb.sheetnames


def test_delete_sheet(excel_file):
    """Test pour supprimer une feuille."""
    delete_sheet(excel_file, "Feuille2")
    wb = openpyxl.load_workbook(excel_file)
    assert "Feuille2" not in wb.sheetnames


def test_set_active_sheet(excel_file):
    """Test pour définir la feuille active."""
    set_active_sheet(excel_file, "Feuille3")
    wb = openpyxl.load_workbook(excel_file)
    assert wb.active.title == "Feuille3"


def test_change_sheet_color(excel_file):
    """Test pour changer la couleur de l'onglet d'une feuille."""
    change_sheet_color(excel_file, "Feuille1", "FF0000")  # Rouge
    wb = openpyxl.load_workbook(excel_file)
    assert wb["Feuille1"].sheet_properties.tabColor == "FF0000"


def test_lock_and_unlock_sheet(excel_file):
    """Test pour verrouiller et déverrouiller une feuille."""
    lock_sheet(excel_file, "Feuille1", password="test123")
    wb = openpyxl.load_workbook(excel_file)
    assert wb["Feuille1"].protection.sheet is True

    unlock_sheet(excel_file, "Feuille1")
    wb = openpyxl.load_workbook(excel_file)
    assert wb["Feuille1"].protection.sheet is False


def test_copy_sheet(excel_file):
    """Test pour copier une feuille."""
    copy_sheet(excel_file, "Feuille1", "Feuille1_Copie")
    wb = openpyxl.load_workbook(excel_file)
    assert "Feuille1_Copie" in wb.sheetnames


def test_reorder_sheets_ascending(excel_file):
    """Test pour réorganiser les feuilles par ordre alphabétique ascendant."""
    reorder_sheets(excel_file, order="asc")
    wb = openpyxl.load_workbook(excel_file)
    assert wb.sheetnames == ["Feuille1", "Feuille2", "Feuille3"]


def test_reorder_sheets_custom(excel_file):
    """Test pour réorganiser les feuilles selon un ordre personnalisé."""
    reorder_sheets(excel_file, order="custom", sheet_order=["Feuille3", "Feuille1", "Feuille2"])
    wb = openpyxl.load_workbook(excel_file)
    assert wb.sheetnames == ["Feuille3", "Feuille1", "Feuille2"]


def test_reorder_sheets_move(excel_file):
    """Test pour déplacer une feuille d'une position à droite."""
    reorder_sheets(excel_file, order="move", sheet_name="Feuille1", direction="right")
    wb = openpyxl.load_workbook(excel_file)
    assert wb.sheetnames == ["Feuille2", "Feuille1", "Feuille3"]


def test_get_sheet_list(excel_file):
    """Test pour obtenir la liste des feuilles."""
    sheets = get_sheet_list(excel_file)
    assert sheets == ["Feuille1", "Feuille2", "Feuille3"]


def test_get_sheet_dict(excel_file):
    """Test pour obtenir un dictionnaire des feuilles avec leurs positions."""
    sheet_dict = get_sheet_dict(excel_file)
    expected_dict = {1: "Feuille1", 2: "Feuille2", 3: "Feuille3"}
    assert sheet_dict == expected_dict
