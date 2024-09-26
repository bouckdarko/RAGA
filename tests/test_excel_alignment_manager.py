import pytest
import openpyxl
from openpyxl.styles import Alignment
from excel_alignment_manager import set_alignment, merge_cells


@pytest.fixture
def excel_file(tmp_path):
    """Crée un fichier Excel temporaire pour les tests."""
    file_path = tmp_path / "test_alignment_workbook.xlsx"
    wb = openpyxl.Workbook()

    # Feuille de test
    sheet = wb.active
    sheet.title = "Feuille1"
    sheet.append(["Col1", "Col2", "Col3"])
    sheet.append([1, 2, 3])
    sheet.append([4, 5, 6])

    wb.save(file_path)
    return file_path


def get_cell_alignment(file_path, sheet_name, cell_ref):
    """Récupère l'alignement d'une cellule."""
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]
    cell = sheet[cell_ref]
    return cell.alignment


def test_set_alignment(excel_file):
    """Test pour appliquer un alignement horizontal et vertical avec renvoi à la ligne."""
    set_alignment(excel_file, "A1:C1", horizontal="center", vertical="bottom", wrap_text=True, sheets=["Feuille1"])
    cell_alignment = get_cell_alignment(excel_file, "Feuille1", "A1")

    # Vérifications sur l'alignement
    assert cell_alignment.horizontal == "center"
    assert cell_alignment.vertical == "bottom"
    assert cell_alignment.wrap_text is True


def test_merge_cells(excel_file):
    """Test pour fusionner une plage de cellules."""
    merge_cells(excel_file, "Feuille1", "A1", "C1")
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb["Feuille1"]

    # Vérification que les cellules A1 à C1 sont fusionnées
    merged_cells = list(sheet.merged_cells.ranges)
    assert f"A1:C1" in str(merged_cells[0])
